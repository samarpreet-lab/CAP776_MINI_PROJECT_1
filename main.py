
from datetime import date
from math import sqrt
from openpyxl import load_workbook


FILE_NAME = __file__.replace("main.py", "12601472.xlsx")
START_ROW = 6                 
START_DATE = date(2026, 8, 13)
END_DATE = date(2026, 9, 20)
EXPECTED_DAYS = (END_DATE - START_DATE).days + 1

FEELING = {"Excellent": 5, "Good": 4, "Neutral": 3, "Low": 2, "Stressed": 1}
SATISFACTION = {
    "Very Satisfied": 5,
    "Satisfied": 4,
    "Neutral": 3,
    "Unsatisfied": 2,
    "Very Unsatisfied": 1,
}
ENERGY = {"High": 3, "Medium": 2, "Low": 1}


def average(values):
    return sum(values) / len(values)


def correlation(x, y):
    x_mean = average(x)
    y_mean = average(y)

    top = sum((a - x_mean) * (b - y_mean) for a, b in zip(x, y))
    bottom_x = sum((a - x_mean) ** 2 for a in x)
    bottom_y = sum((b - y_mean) ** 2 for b in y)

    return top / sqrt(bottom_x * bottom_y)


# Open the Excel file and use the first sheet (Daily Log).
workbook = load_workbook(FILE_NAME, data_only=True)
sheet = workbook.active

sleep = []
fitness = []
study = []
coding = []
class_time = []
other = []
total_tracked = []
free_time = []
feeling_score = []
satisfaction_score = []
energy_score = []


for row in sheet.iter_rows(min_row=START_ROW, values_only=True):
    if row[0] is None:       # Ignore empty rows.
        continue

    sleep_value = float(row[1])
    fitness_value = float(row[2])
    study_value = float(row[3])
    coding_value = float(row[4])
    class_value = float(row[5])
    other_value = float(row[7])

    # Total and free time are calculated from the activity values.
    total_value = sleep_value + fitness_value + study_value + coding_value + class_value + other_value
    free_value = 1440 - total_value

    sleep.append(sleep_value)
    fitness.append(fitness_value)
    study.append(study_value)
    coding.append(coding_value)
    class_time.append(class_value)
    other.append(other_value)
    total_tracked.append(total_value)
    free_time.append(free_value)
    feeling_score.append(FEELING[row[10]])
    satisfaction_score.append(SATISFACTION[row[11]])
    energy_score.append(ENERGY[row[12]])


# Activity Data Summary
valid_days = len(sleep)
missing_days = EXPECTED_DAYS - valid_days
invalid_records = 0

avg_sleep = average(sleep)
avg_fitness = average(fitness)
avg_study = average(study)
avg_coding = average(coding)
avg_class = average(class_time)
avg_other = average(other)
avg_free = average(free_time)

# Index Values
TPI = avg_coding
AAI = average([study[i] + class_time[i] for i in range(valid_days)])
PhAI = avg_fitness
SRI = avg_sleep
ABI = avg_free
TUI = average(total_tracked)
EI = average([
    (feeling_score[i] + satisfaction_score[i] + energy_score[i]) / 3
    for i in range(valid_days)
])
DCI = (valid_days / EXPECTED_DAYS) * 100
PAI = (0.15 * TPI) + (0.20 * AAI) + (0.15 * PhAI) + (0.20 * SRI) + \
      (0.15 * TUI) + (0.10 * EI) + (0.05 * DCI)

# Required correlations
sleep_energy = correlation(sleep, energy_score)
study_satisfaction = correlation(study, satisfaction_score)
coding_energy = correlation(coding, energy_score)


print("CAP776 MINOR PROJECT #1 - MY DATA, MY STORY")
print("\n1. ACTIVITY DATA SUMMARY")
print("Expected number of days:", EXPECTED_DAYS)
print("Valid days recorded:", valid_days)
print("Missing days:", missing_days)
print("Invalid / excluded records:", invalid_records)
print(f"Average Sleep/day: {avg_sleep:.2f} min/day")
print(f"Average Fitness/day: {avg_fitness:.2f} min/day")
print(f"Average Study/day: {avg_study:.2f} min/day")
print(f"Average Coding/day: {avg_coding:.2f} min/day")
print(f"Average Class/day: {avg_class:.2f} min/day")
print(f"Average Other Activities/day: {avg_other:.2f} min/day")
print(f"Average Free / Unaccounted Time/day: {avg_free:.2f} min/day")

print("\n2. INDEX VALUES")
print(f"PAI: {PAI:.2f}")
print(f"TPI: {TPI:.2f} min/day")
print(f"AAI: {AAI:.2f} min/day")
print(f"PhAI: {PhAI:.2f} min/day")
print(f"SRI: {SRI:.2f} min/day")
print(f"ABI: {ABI:.2f} min/day")
print(f"TUI: {TUI:.2f} min/day")
print(f"EI: {EI:.2f} / 5")
print(f"DCI: {DCI:.2f}%")

print("\n3. REQUIRED CORRELATIONS")
print(f"Sleep - Energy: {sleep_energy:.2f}")
print(f"Study - Satisfaction: {study_satisfaction:.2f}")
print(f"Coding - Energy: {coding_energy:.2f}")
